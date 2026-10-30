import tempfile
import unittest
from pathlib import Path

try:
    from openpyxl import Workbook
    from ingestion.catalog_v2 import CatalogV2ParseError, REQUIRED_SHEETS, parse_catalog_v2
except ModuleNotFoundError:  # pragma: no cover - environment dependency gate
    Workbook = None


class CatalogV2Tests(unittest.TestCase):
    def setUp(self) -> None:
        if Workbook is None:
            self.skipTest("openpyxl is not installed")

    def workbook(self, directory: str) -> Path:
        path = Path(directory) / "catalog-v2.xlsx"
        workbook = Workbook()
        workbook.active.title = "README"
        headers = {
            "README": ["key", "value"],
            "CatalogManifest": ["schema_version", "catalog_key", "tenant_scope", "revision", "policy_version"],
            "BusinessGoals": ["business_goal_key", "name", "description", "active", "confirmation_required", "completion_policy_key"],
            "IntentCategories": ["category_key", "name", "active"],
            "TurnIntents": ["turn_intent_key", "category_key", "business_goal_key", "active"],
            "IntentSignals": ["signal_id", "target_key", "signal_type", "value", "language", "weight", "min_weight", "max_weight"],
            "ScoringPolicies": ["scoring_policy_key", "business_goal_key", "threshold", "ambiguity_margin", "policy_version"],
            "RequirementPolicies": ["requirement_key", "goal_key", "stage_key", "type", "required", "dependencies", "validation", "source_of_truth", "prompt_key", "priority"],
            "BusinessStages": ["business_goal_key", "stage_key", "stage_order", "entry_requirements", "completion_evidence_policy_key", "active"],
            "ToolBindings": ["tool_key", "goal_key", "stage_key", "side_effect_class", "arguments_schema_ref", "required_keys", "consent", "confirmation", "registry_key", "enabled"],
            "WorkflowBindings": ["workflow_key", "registry_key", "provider_type", "payload_schema_ref", "callback_schema_ref", "evidence_policy_key", "enabled"],
            "FollowUpPolicies": ["policy_key", "goal_key", "stage_key", "requirement_key", "priority", "no_question_reasons"],
            "KnowledgeTopics": ["topic_key", "goal_key", "source_scope", "citation_required", "active"],
            "PersonaPolicies": ["persona_policy_key", "persona_version_id", "permitted_languages", "fallback_template_ref", "enabled"],
            "DocumentPolicies": ["policy_key", "goal_key", "stage_key", "doc_type", "required", "mime_types", "review_policy", "retention_class", "enabled"],
        }
        for name in REQUIRED_SHEETS:
            sheet = workbook[name] if name == "README" else workbook.create_sheet(name)
            sheet.append(headers[name])
        workbook["CatalogManifest"].append([2, "ancfz", "ancfz", 1, "ancfz-v2"])
        workbook.save(path)
        return path

    def test_valid_workbook_is_hash_bound_and_lossless(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            parsed = parse_catalog_v2(self.workbook(directory), "ancfz")
            self.assertEqual(parsed["schemaVersion"], 2)
            self.assertEqual(parsed["manifest"]["catalogKey"], "ancfz")
            self.assertEqual(len(parsed["sourceHash"]), 64)

    def test_rejects_v1_unknown_sheet_and_scope_mismatch(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = self.workbook(directory)
            workbook = __import__("openpyxl").load_workbook(path)
            workbook["CatalogManifest"]["A2"] = 1
            workbook.create_sheet("Intents")
            workbook.save(path)
            with self.assertRaises(CatalogV2ParseError):
                parse_catalog_v2(path)
            self.workbook(directory)
            with self.assertRaises(CatalogV2ParseError) as error:
                parse_catalog_v2(path, "other-tenant")
            self.assertEqual(error.exception.code, "TENANT_SCOPE_MISMATCH")

    def test_rejects_formula_and_bad_json_array(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = self.workbook(directory)
            workbook = __import__("openpyxl").load_workbook(path)
            workbook["BusinessStages"].append(["ancfz", "lead_generation", 1, "not-json", "evidence", True])
            workbook.save(path)
            with self.assertRaises(CatalogV2ParseError) as error:
                parse_catalog_v2(path)
            self.assertIn(error.exception.code, {"INVALID_JSON_ARRAY", "FORMULA_NOT_ALLOWED"})
            self.workbook(directory)
            workbook = __import__("openpyxl").load_workbook(path)
            workbook["BusinessStages"]["D2"] = "=SUM(1,1)"
            workbook.save(path)
            with self.assertRaises(CatalogV2ParseError) as error:
                parse_catalog_v2(path)
            self.assertEqual(error.exception.code, "FORMULA_NOT_ALLOWED")


if __name__ == "__main__":
    unittest.main()
