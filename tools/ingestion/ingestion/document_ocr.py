from __future__ import annotations

import io
import re
from pathlib import Path


def _clean(value: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", value.replace("\x00", " "))).strip()


def extract_document_text(file_name: str, mime_type: str, content: bytes, languages: list[str] | None = None, max_pdf_pages: int = 20) -> str:
    languages = languages or ["eng", "ara"]
    if mime_type in {"text/plain", "text/csv"}:
        return _clean(content.decode("utf-8-sig", errors="replace"))
    if mime_type == "application/pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        native = _clean("\n".join(page.extract_text() or "" for page in reader.pages[:max_pdf_pages]))
        if len(native) >= 40:
            return native
        import fitz
        from PIL import Image
        import pytesseract
        document = fitz.open(stream=content, filetype="pdf")
        pages = []
        for page in document[:max_pdf_pages]:
            pixels = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            image = Image.frombytes("RGB", [pixels.width, pixels.height], pixels.samples)
            pages.append(pytesseract.image_to_string(image, lang="+".join(languages)))
        return _clean("\n".join(pages))
    if mime_type.startswith("image/"):
        from PIL import Image
        import pytesseract
        return _clean(pytesseract.image_to_string(Image.open(io.BytesIO(content)), lang="+".join(languages)))
    if mime_type.endswith("wordprocessingml.document"):
        from docx import Document
        document = Document(io.BytesIO(content))
        return _clean("\n".join(paragraph.text for paragraph in document.paragraphs))
    if mime_type.endswith("spreadsheetml.sheet"):
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return _clean("\n".join(" | ".join(str(cell.value or "") for cell in row) for sheet in workbook.worksheets for row in sheet.iter_rows()))
    raise ValueError(f"Unsupported document type: {mime_type}")


def document_summary(text: str, file_name: str) -> str:
    excerpt = _clean(text)[:4000]
    return f"OCR extracted business-document text from {Path(file_name).name}: {excerpt}" if excerpt else f"OCR completed for {Path(file_name).name}, but no readable text was found."
