"""Strict, v2-only ANCFZ catalog workbook parser.

This module is deliberately separate from intent_catalog.py, which remains the
legacy v1 parser used for historical catalog records and rollback.
"""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook

PARSER_VERSION = "intent-catalog-v2-xlsx-1.0.0"
REQUIRED_SHEETS = (
    "README", "CatalogManifest", "BusinessGoals", "IntentCategories", "TurnIntents",
    "IntentSignals", "ScoringPolicies", "RequirementPolicies", "BusinessStages",
    "ToolBindings", "WorkflowBindings", "FollowUpPolicies", "KnowledgeTopics",
    "PersonaPolicies", "DocumentPolicies",
)

HEADERS: dict[str, tuple[str, ...]] = {
    "README": ("key", "value"),
    "CatalogManifest": ("schema_version", "catalog_key", "tenant_scope", "revision", "policy_version"),
    "BusinessGoals": ("business_goal_key", "name", "description", "active", "confirmation_required", "completion_policy_key"),
    "IntentCategories": ("category_key", "name", "active"),
    "TurnIntents": ("turn_intent_key", "category_key", "business_goal_key", "active"),
    "IntentSignals": ("signal_id", "target_key", "signal_type", "value", "language", "weight", "min_weight", "max_weight"),
    "ScoringPolicies": ("scoring_policy_key", "business_goal_key", "threshold", "ambiguity_margin", "policy_version"),
    "RequirementPolicies": ("requirement_key", "goal_key", "stage_key", "type", "required", "dependencies", "validation", "source_of_truth", "prompt_key", "priority"),
    "BusinessStages": ("business_goal_key", "stage_key", "stage_order", "entry_requirements", "completion_evidence_policy_key", "active"),
    "ToolBindings": ("tool_key", "goal_key", "stage_key", "side_effect_class", "arguments_schema_ref", "required_keys", "consent", "confirmation", "registry_key", "enabled"),
    "WorkflowBindings": ("workflow_key", "registry_key", "provider_type", "payload_schema_ref", "callback_schema_ref", "evidence_policy_key", "enabled"),
    "FollowUpPolicies": ("policy_key", "goal_key", "stage_key", "requirement_key", "priority", "no_question_reasons"),
    "KnowledgeTopics": ("topic_key", "goal_key", "source_scope", "citation_required", "active"),
    "PersonaPolicies": ("persona_policy_key", "persona_version_id", "permitted_languages", "fallback_template_ref", "enabled"),
    "DocumentPolicies": ("policy_key", "goal_key", "stage_key", "doc_type", "required", "mime_types", "review_policy", "retention_class", "enabled"),
}

ARRAY_FIELDS = {"dependencies", "entry_requirements", "required_keys", "no_question_reasons", "permitted_languages", "mime_types"}
BOOL_FIELDS = {"active", "confirmation_required", "required", "enabled", "citation_required"}
INT_FIELDS = {"schema_version", "revision", "stage_order", "priority"}
FLOAT_FIELDS = {"threshold", "ambiguity_margin", "weight", "min_weight", "max_weight"}
CAMEL = re.compile(r"_([a-z])")


class CatalogV2ParseError(ValueError):
    def __init__(self, code: str, message: str, sheet: str | None = None, row: int | None = None, column: str | None = None) -> None:
        self.code, self.sheet, self.row, self.column = code, sheet, row, column
        location = ".".join(part for part in (sheet, str(row) if row is not None else None, column) if part)
        super().__init__(f"{code}: {message}{f' [{location}]' if location else ''}")


def _camel(value: str) -> str:
    return CAMEL.sub(lambda match: match.group(1).upper(), value)


def _cell_value(value: Any, field: str, sheet: str, row: int, column: str) -> Any:
    if value is None or (isinstance(value, str) and not value.strip()):
        return [] if field in ARRAY_FIELDS else None
    if field in ARRAY_FIELDS:
        if not isinstance(value, str):
            raise CatalogV2ParseError("INVALID_JSON_ARRAY", "Array fields must contain JSON arrays.", sheet, row, column)
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError as error:
            raise CatalogV2ParseError("INVALID_JSON_ARRAY", "Array field is not valid JSON.", sheet, row, column) from error
        if not isinstance(parsed, list):
            raise CatalogV2ParseError("INVALID_JSON_ARRAY", "Array field must decode to a JSON array.", sheet, row, column)
        return parsed
    if field in BOOL_FIELDS:
        if isinstance(value, bool):
            return value
        if isinstance(value, str) and value.strip().lower() in {"true", "false"}:
            return value.strip().lower() == "true"
        raise CatalogV2ParseError("INVALID_BOOLEAN", "Expected true or false.", sheet, row, column)
    if field in INT_FIELDS:
        if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
            raise CatalogV2ParseError("INVALID_INTEGER", "Expected an integer.", sheet, row, column)
        return int(value)
    if field in FLOAT_FIELDS:
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise CatalogV2ParseError("INVALID_NUMBER", "Expected a number.", sheet, row, column)
        return float(value)
    return str(value).strip()


def _sheet_rows(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows())
    if not rows:
        raise CatalogV2ParseError("MISSING_HEADERS", "Every required sheet must contain a header row.", sheet_name, 1)
    header_values = [str(cell.value).strip() if cell.value is not None else "" for cell in rows[0]]
    expected = list(HEADERS[sheet_name])
    if header_values != expected:
        raise CatalogV2ParseError("INVALID_HEADERS", f"Expected headers {expected!r}.", sheet_name, 1)
    result: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        values = [cell.value for cell in row[:len(expected)]]
        if not any(value is not None and str(value).strip() for value in values):
            continue
        for cell in row:
            if cell.data_type == "f":
                raise CatalogV2ParseError("FORMULA_NOT_ALLOWED", "Workbook formulas are not allowed.", sheet_name, row_number, cell.column_letter)
        result.append({
            _camel(field): _cell_value(value, field, sheet_name, row_number, column)
            for field, value, column in zip(expected, values, (cell.column_letter for cell in row[:len(expected)]), strict=True)
        })
    return result


def parse_catalog_v2(path: Path, expected_tenant_scope: str | None = None) -> dict[str, Any]:
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise CatalogV2ParseError("INVALID_FILE", "An existing .xlsx workbook is required.")
    source = path.read_bytes()
    with ZipFile(path) as archive:
        if any(name.startswith("xl/externalLinks/") or name.endswith("vbaProject.bin") for name in archive.namelist()):
            raise CatalogV2ParseError("UNSAFE_WORKBOOK", "External links and macros are not allowed.")
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        unknown = sorted(set(workbook.sheetnames) - set(REQUIRED_SHEETS))
        missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
        if missing:
            raise CatalogV2ParseError("MISSING_SHEETS", f"Missing required sheets: {', '.join(missing)}")
        if unknown:
            raise CatalogV2ParseError("UNKNOWN_SHEET", f"Unknown executable sheets: {', '.join(unknown)}")
        parsed = {name: _sheet_rows(workbook, name) for name in REQUIRED_SHEETS}
        if len(parsed["CatalogManifest"]) != 1:
            raise CatalogV2ParseError("INVALID_MANIFEST", "CatalogManifest must contain exactly one row.", "CatalogManifest")
        manifest = parsed["CatalogManifest"][0]
        if manifest.get("schemaVersion") != 2:
            raise CatalogV2ParseError("UNSUPPORTED_SCHEMA", "Only schemaVersion=2 is accepted.", "CatalogManifest", 2, "A")
        if expected_tenant_scope is not None and manifest.get("tenantScope") != expected_tenant_scope:
            raise CatalogV2ParseError("TENANT_SCOPE_MISMATCH", "Workbook tenant scope does not match the target tenant.", "CatalogManifest", 2, "C")
        return {
            "schemaVersion": 2,
            "manifest": manifest,
            "sheets": parsed,
            "sourceHash": hashlib.sha256(source).hexdigest(),
            "parserVersion": PARSER_VERSION,
        }
    finally:
        workbook.close()
