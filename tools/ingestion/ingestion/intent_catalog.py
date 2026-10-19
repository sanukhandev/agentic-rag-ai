from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PARSER_VERSION = "intent-catalog-xlsx-1.0.0"
SHEETS = ("README", "Intents", "Signals", "Slots", "Relations", "WorkflowBindings")


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _camel(value: str) -> str:
    parts = value.replace(" ", "_").replace("-", "_").split("_")
    return parts[0][:1].lower() + parts[0][1:] + "".join(part[:1].upper() + part[1:] for part in parts[1:])


def _rows(sheet: Any) -> list[dict[str, Any]]:
    values = list(sheet.values)
    if not values:
        return []
    headers = [_camel(_text(value)) for value in values[0]]
    if any(not header for header in headers) or len(set(headers)) != len(headers):
        raise ValueError(f"Sheet {sheet.title} has invalid or duplicate headers.")
    return [dict(zip(headers, row, strict=False)) for row in values[1:] if any(value is not None and _text(value) for value in row)]


def _readme_metadata(sheet: Any) -> dict[str, str]:
    return {
        _camel(_text(row[0])): _text(row[1])
        for row in sheet.values
        if len(row) >= 2 and _text(row[0]) and _text(row[1])
    }


def _reject_formulas(workbook: Any) -> None:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError(f"Workbook formula is not allowed: {sheet.title}!{cell.coordinate}")
    if getattr(workbook, "_external_links", []):
        raise ValueError("Workbook external links are not allowed.")


def parse_intent_catalog(path: Path) -> dict[str, Any]:
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise ValueError("An existing .xlsx workbook is required.")
    content = path.read_bytes()
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        _reject_formulas(workbook)
        missing = [sheet for sheet in SHEETS if sheet not in workbook.sheetnames]
        if missing:
            raise ValueError(f"Workbook is missing required sheets: {', '.join(missing)}")
        tenant_scope = _readme_metadata(workbook["README"]).get("tenantScope")
        return {
            "sourceHash": hashlib.sha256(content).hexdigest(),
            "parserVersion": PARSER_VERSION,
            "catalog": {
                "tenantScope": tenant_scope or None,
                "intents": _rows(workbook["Intents"]),
                "signals": _rows(workbook["Signals"]),
                "slots": _rows(workbook["Slots"]),
                "relations": _rows(workbook["Relations"]),
                "workflowBindings": _rows(workbook["WorkflowBindings"]),
            },
        }
    finally:
        workbook.close()


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Validate and convert an Excel intent catalog to JSON.")
    parser.add_argument("path", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.write_text(json.dumps(parse_intent_catalog(args.path), ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")


if __name__ == "__main__":
    main()
