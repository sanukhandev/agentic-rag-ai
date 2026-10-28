from __future__ import annotations

import csv
import hashlib
import ipaddress
import json
import re
import socket
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from pathlib import PurePosixPath
from urllib.parse import urlparse

MAX_BYTES = 20 * 1024 * 1024
MAX_CONTAINER_ENTRIES = 2_000
MAX_CONTAINER_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_NORMALIZED_BYTES = 25 * 1024 * 1024
MAX_RESULT_CHUNKS = 100_000
PARSER_VERSION = "python-ingestion-0.4.0"
ARCHIVE_SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".csv", ".json", ".jsonl"}


@dataclass(frozen=True)
class Chunk:
    ordinal: int
    locator: str
    content: str
    content_hash: str


@dataclass(frozen=True)
class IngestedSource:
    title: str
    mime_type: str
    source_url: str | None
    normalized_content: str
    content_hash: str
    parser_version: str
    chunks: list[Chunk]
    warnings: list[str] = field(default_factory=list)


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def sanitize(text: str) -> str:
    text = re.sub(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", "[REDACTED_EMAIL]", text)
    text = re.sub(r"(?<!\w)(?:\+?\d[\d ()-]{7,}\d)(?!\w)", "[REDACTED_PHONE]", text)
    text = re.sub(r"(?i)(api[_ -]?key|secret|password|token)\s*[:=]\s*[^\s,;]+", r"\1=[REDACTED_SECRET]", text)
    return text


def chunk_text(text: str, max_chars: int = 800) -> list[Chunk]:
    paragraphs = [part.strip() for part in normalize(text).split("\n\n") if part.strip()]
    output: list[Chunk] = []
    current = ""
    for paragraph in paragraphs:
        if current and len(current) + len(paragraph) + 2 > max_chars:
            output.append(_chunk(current, len(output)))
            current = ""
        current = f"{current}\n\n{paragraph}".strip()
        if len(current) > max_chars:
            for start in range(0, len(current), max_chars):
                piece = current[start : start + max_chars].strip()
                if piece:
                    output.append(_chunk(piece, len(output)))
            current = ""
    if current:
        output.append(_chunk(current, len(output)))
    return output


def _chunk(content: str, ordinal: int, locator: str | None = None) -> Chunk:
    return Chunk(ordinal=ordinal, locator=locator or f"section {ordinal + 1}", content=content, content_hash=hashlib.sha256(content.encode()).hexdigest())


def _safe_path(path: Path) -> Path:
    resolved = path.resolve()
    if not resolved.is_file():
        raise ValueError("Input file does not exist.")
    if resolved.stat().st_size > MAX_BYTES:
        raise ValueError("Input exceeds the 20 MiB ingestion limit.")
    return resolved


def _validate_file_container(path: Path) -> None:
    suffix = path.suffix.lower()
    with path.open("rb") as handle:
        header = handle.read(5)
    if suffix == ".pdf" and header != b"%PDF-":
        raise ValueError("File signature does not match PDF.")
    if suffix not in {".docx", ".xlsx"}:
        return
    if not zipfile.is_zipfile(path):
        raise ValueError("File signature does not match Office Open XML.")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_CONTAINER_ENTRIES or sum(entry.file_size for entry in entries) > MAX_CONTAINER_UNCOMPRESSED_BYTES:
            raise ValueError("Office container exceeds safe extraction limits.")
        if any(PurePosixPath(entry.filename).is_absolute() or ".." in PurePosixPath(entry.filename).parts for entry in entries):
            raise ValueError("Office container contains an unsafe path.")
        names = {entry.filename.lower() for entry in entries}
        if any(name.endswith("vbaproject.bin") for name in names):
            raise ValueError("Macro-enabled Office files are unsupported.")


def _public_addresses(hostname: str, port: int) -> set[ipaddress.IPv4Address | ipaddress.IPv6Address]:
    try:
        addresses = {ipaddress.ip_address(info[4][0]) for info in socket.getaddrinfo(hostname, port, type=socket.SOCK_STREAM)}
    except (OSError, ValueError) as error:
        raise ValueError("Website host could not be resolved safely.") from error
    if any(address.is_private or address.is_loopback or address.is_link_local or address.is_reserved or address.is_multicast or address.is_unspecified for address in addresses):
        raise ValueError("Private and local network targets are not allowed.")
    return addresses


def _response_peer(response: object) -> ipaddress.IPv4Address | ipaddress.IPv6Address | None:
    raw = getattr(response, "raw", None)
    connection = getattr(raw, "_connection", None) or getattr(raw, "connection", None)
    sock = getattr(connection, "sock", None)
    if sock is None:
        return None
    try:
        return ipaddress.ip_address(sock.getpeername()[0])
    except (OSError, ValueError, IndexError, TypeError):
        return None


def extract_file(path: Path) -> tuple[str, str]:
    path = _safe_path(path)
    _validate_file_container(path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8"), "text/markdown" if suffix == ".md" else "text/plain"
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8") as handle:
            return "\n".join(" | ".join(row) for row in csv.reader(handle)), "text/csv"
    if suffix == ".json":
        try:
            value = json.loads(path.read_text(encoding="utf-8-sig"))
        except json.JSONDecodeError as error:
            raise ValueError("Invalid JSON source.") from error
        return json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True), "application/json"
    if suffix == ".jsonl":
        records: list[str] = []
        for line_number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), start=1):
            if not line.strip():
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as error:
                raise ValueError(f"Invalid JSONL source at line {line_number}.") from error
            records.append(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True))
        if not records:
            raise ValueError("JSONL source contains no records.")
        return "\n\n".join(records), "application/x-ndjson"
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        parts = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = str(getattr(paragraph.style, "name", ""))
            heading = re.fullmatch(r"Heading\s+(\d+)", style_name, re.IGNORECASE)
            parts.append(f"Heading {heading.group(1)}: {text}" if heading else text)
        parts.extend(
            f"Table {table_index + 1}, row {row_index + 1}: " + " | ".join(cell.text.strip() for cell in row.cells)
            for table_index, table in enumerate(document.tables)
            for row_index, row in enumerate(table.rows)
        )
        return "\n\n".join(parts), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = []
            for sheet in workbook.worksheets:
                values.append(f"Sheet: {sheet.title}")
                values.extend(" | ".join("" if cell is None else str(cell) for cell in row) for row in sheet.iter_rows(values_only=True))
            return "\n".join(values), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        finally:
            workbook.close()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        pages = [f"Page {index}:\n{text}" for index, page in enumerate(reader.pages, start=1) if (text := page.extract_text() or "").strip()]
        if not pages:
            raise ValueError("PDF has no extractable text; OCR-only or empty PDFs are unsupported.")
        return "\n\n".join(pages), "application/pdf"
    if suffix == ".zip":
        raise ValueError("ZIP archives must be processed through ingest_file.")
    raise ValueError(f"Unsupported file type: {suffix}")


def _is_symlink(entry: zipfile.ZipInfo) -> bool:
    return ((entry.external_attr >> 16) & 0o170000) == 0o120000


def _archive_member_label(name: str) -> str:
    return name.replace("\r", " ").replace("\n", " ")[:200]


def _extract_archive(path: Path) -> tuple[list[tuple[str, str, str]], list[str]]:
    if not zipfile.is_zipfile(path):
        raise ValueError("File signature does not match ZIP archive.")
    with zipfile.ZipFile(path) as archive:
        entries = [entry for entry in archive.infolist() if not entry.is_dir()]
        if not entries:
            raise ValueError("ZIP archive contains no files.")
        if len(entries) > MAX_CONTAINER_ENTRIES or sum(entry.file_size for entry in entries) > MAX_CONTAINER_UNCOMPRESSED_BYTES:
            raise ValueError("ZIP archive exceeds safe extraction limits.")
        seen_names: set[str] = set()
        for entry in entries:
            name = entry.filename
            relative = PurePosixPath(name)
            if (
                not name
                or "\\" in name
                or "\x00" in name
                or relative.is_absolute()
                or ".." in relative.parts
                or _is_symlink(entry)
            ):
                raise ValueError("ZIP archive contains an unsafe path.")
            if name.casefold() in seen_names:
                raise ValueError("ZIP archive contains duplicate file names.")
            seen_names.add(name.casefold())
            if entry.flag_bits & 0x1 or entry.file_size > MAX_BYTES:
                raise ValueError("ZIP archive contains an encrypted or oversized file.")
        has_training_directory = any("training" in {part.casefold() for part in PurePosixPath(entry.filename).parts[:-1]} for entry in entries)
        if has_training_directory:
            entries = [entry for entry in entries if "training" in {part.casefold() for part in PurePosixPath(entry.filename).parts[:-1]}]
        if not entries:
            raise ValueError("ZIP archive contains no ingestible files.")
        extracted: list[tuple[str, str, str]] = []
        warnings: list[str] = []
        with tempfile.TemporaryDirectory(prefix="agentic-zip-") as directory:
            root = Path(directory)
            for index, entry in enumerate(entries):
                name = entry.filename
                relative = PurePosixPath(name)
                suffix = Path(name).suffix.lower()
                if suffix == ".zip":
                    warnings.append(f"Skipped {_archive_member_label(name)}: nested ZIP archives are unsupported.")
                    continue
                if suffix not in ARCHIVE_SUPPORTED_SUFFIXES:
                    warnings.append(f"Skipped {_archive_member_label(name)}: unsupported file type {suffix or '[none]'}.")
                    continue
                if entry.flag_bits & 0x1 or entry.file_size > MAX_BYTES:
                    raise ValueError("ZIP archive contains an encrypted or oversized file.")
                try:
                    with archive.open(entry) as source:
                        content = source.read(MAX_BYTES + 1)
                    if len(content) != entry.file_size or len(content) > MAX_BYTES:
                        raise ValueError("oversized or truncated file")
                    member = root / f"{index:04d}-{relative.name}"
                    member.write_bytes(content)
                    text, mime_type = extract_file(member)
                    if not normalize(text):
                        raise ValueError("empty file")
                except ValueError as error:
                    warnings.append(f"Skipped {_archive_member_label(name)}: {error}.")
                    continue
                extracted.append((name, mime_type, text))
        if not extracted:
            detail = " ".join(warnings[:10])
            raise ValueError(f"ZIP archive contains no supported files.{f' {detail}' if detail else ''}")
        return extracted, warnings


def _ingest_archive(path: Path) -> IngestedSource:
    files, warnings = _extract_archive(path)
    sections: list[str] = []
    chunks: list[Chunk] = []
    for name, mime_type, text in files:
        display_name = sanitize(name)
        cleaned = sanitize(normalize(text))
        sections.append(f"File: {display_name}\nFormat: {mime_type}\n{cleaned}")
        for chunk in chunk_text(cleaned):
            content = f"File: {display_name}\n{chunk.content}"
            chunks.append(_chunk(content, len(chunks), f"{display_name} · {chunk.locator}"))
    normalized_content = normalize("\n\n".join(sections))
    if len(normalized_content.encode("utf-8")) > MAX_NORMALIZED_BYTES or len(chunks) > MAX_RESULT_CHUNKS:
        raise ValueError("Normalized ZIP ingestion result exceeds the configured output budget; split the archive into smaller files.")
    return IngestedSource(
        title=path.stem,
        mime_type="application/zip",
        source_url=None,
        normalized_content=normalized_content,
        content_hash=hashlib.sha256(normalized_content.encode()).hexdigest(),
        parser_version=PARSER_VERSION,
        chunks=chunks,
        warnings=warnings,
    )


def fetch_authorized_https(url: str, allowed_hosts: set[str]) -> tuple[str, str]:
    parsed = urlparse(url)
    hostname = parsed.hostname.lower() if parsed.hostname else ""
    allowed = {host.strip().lower() for host in allowed_hosts}
    if parsed.scheme != "https" or not hostname or hostname not in allowed or parsed.username or parsed.password or parsed.port not in (None, 443):
        raise ValueError("Website ingestion requires an allowlisted HTTPS host.")
    addresses = _public_addresses(hostname, parsed.port or 443)
    import requests
    from bs4 import BeautifulSoup

    response = requests.get(url, timeout=10, headers={"User-Agent": "agentic-rag-ingestion/0.1"}, allow_redirects=False, stream=True)
    response.raise_for_status()
    if 300 <= response.status_code < 400:
        raise ValueError("Redirects are not accepted for website ingestion; approve the final host explicitly.")
    content_type = response.headers.get("content-type", "").split(";", 1)[0].strip().lower()
    if content_type not in {"text/html", "application/xhtml+xml", "text/plain"}:
        raise ValueError("Website response must be text content.")
    rebound_addresses = _public_addresses(hostname, parsed.port or 443)
    peer = _response_peer(response)
    if peer and (peer.is_private or peer.is_loopback or peer.is_link_local or peer.is_reserved or peer.is_multicast or peer.is_unspecified or peer not in addresses | rebound_addresses):
        raise ValueError("Website connection did not remain on an approved public address.")
    chunks: list[bytes] = []
    size = 0
    for chunk in response.iter_content(chunk_size=64 * 1024):
        size += len(chunk)
        if size > MAX_BYTES:
            raise ValueError("Website response exceeds the 20 MiB ingestion limit.")
        chunks.append(chunk)
    soup = BeautifulSoup(b"".join(chunks).decode(response.encoding or "utf-8", errors="replace"), "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()
    return soup.get_text("\n"), "text/html"


def ingest_text(title: str, text: str, mime_type: str, source_url: str | None = None) -> IngestedSource:
    cleaned = sanitize(normalize(text))
    chunks = chunk_text(cleaned)
    if len(cleaned.encode("utf-8")) > MAX_NORMALIZED_BYTES or len(chunks) > MAX_RESULT_CHUNKS:
        raise ValueError("Normalized ingestion result exceeds the configured output budget; split the source into smaller files.")
    return IngestedSource(title=title, mime_type=mime_type, source_url=source_url, normalized_content=cleaned, content_hash=hashlib.sha256(cleaned.encode()).hexdigest(), parser_version=PARSER_VERSION, chunks=chunks)


def ingest_file(path: Path) -> IngestedSource:
    path = _safe_path(path)
    if path.suffix.lower() == ".zip":
        return _ingest_archive(path)
    text, mime_type = extract_file(path)
    return ingest_text(path.stem, text, mime_type)


def ingest_url(title: str, url: str, allowed_hosts: set[str]) -> IngestedSource:
    text, mime_type = fetch_authorized_https(url, allowed_hosts)
    return ingest_text(title, text, mime_type, url)


def to_json(source: IngestedSource) -> str:
    return json.dumps(asdict(source), indent=2, sort_keys=True)
