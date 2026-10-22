import tempfile
import unittest
import zipfile
import sys
import requests
from unittest.mock import Mock, patch
from pathlib import Path

from ingestion.pipeline import chunk_text, extract_file, fetch_authorized_https, ingest_file, ingest_text, ingest_url, sanitize
from ingestion.intent_catalog import parse_intent_catalog
from ingestion.staging import LocalStaging
from ingestion.worker import claim_next_job, fetch_job_input, main as worker_main, renew_document_job, result_payload, submit_failure, submit_result


class IngestionTests(unittest.TestCase):
    def test_intent_catalog_workbook_is_structured_and_deterministic(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "catalog.xlsx"
            workbook = Workbook()
            workbook.active.title = "README"
            workbook["README"].append(["Tenant scope", "demo"])
            for name, headers, row in [
                ("Intents", ["intent_key", "display_name", "workflow_key", "minimum_score", "confidence_threshold", "confirmation_required", "active", "priority"], ["lead.capture", "Lead", "crm.lead.create", 70, 0.8, True, True, 100]),
                ("Signals", ["intent_key", "signal_id", "signal_type", "value", "weight"], ["lead.capture", "lead-1", "phrase", "need a demo", 40]),
                ("Slots", ["intent_key", "slot_key", "data_type", "required", "collection_order", "prompt_key", "validation_rule"], ["lead.capture", "name", "person_name", True, 1, "lead.ask_name", "token:NAME"]),
                ("Relations", ["source_intent", "related_intent", "relation_type", "trigger_group", "score_adjustment", "minimum_cooccurrence", "clarification_prompt_key"], ["lead.capture", "lead.capture", "refines", "demo", 0, 1, "clarify"]),
                ("WorkflowBindings", ["workflow_key", "n8n_workflow_id", "webhook_path", "payload_contract", "side_effect_class", "active"], ["crm.lead.create", "demo", "lead-capture", "crm-lead-v1", "crm_write", True]),
            ]:
                sheet = workbook.create_sheet(name)
                sheet.append(headers)
                sheet.append(row)
            workbook.save(path)
            first = parse_intent_catalog(path)
            second = parse_intent_catalog(path)
            self.assertEqual(first, second)
            self.assertEqual(first["catalog"]["intents"][0]["intentKey"], "lead.capture")
            self.assertEqual(first["catalog"]["tenantScope"], "demo")
            self.assertEqual(first["catalog"]["qualityRules"], [])
            self.assertEqual(first["catalog"]["catalogManifest"], [])
    def test_masks_pii_and_secrets_before_chunking(self) -> None:
        result = sanitize("Email alice@example.com, call +971 50 123 4567, api_key=secret-value")
        self.assertNotIn("alice@example.com", result)
        self.assertNotIn("secret-value", result)
        self.assertIn("[REDACTED_EMAIL]", result)

    def test_supported_text_file_is_deterministic(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "faq.txt"
            path.write_text("Pricing is published.\n\nSupport is available.", encoding="utf-8")
            first = ingest_file(path)
            second = ingest_file(path)
            self.assertEqual(first.content_hash, second.content_hash)
            self.assertEqual(first.normalized_content, "Pricing is published.\n\nSupport is available.")
            self.assertEqual(first.chunks, second.chunks)

    def test_chunk_bound_is_enforced(self) -> None:
        self.assertTrue(all(len(chunk.content) <= 20 for chunk in chunk_text("one two three four five six", max_chars=20)))

    def test_unsupported_file_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "file.exe"
            path.write_bytes(b"not allowed")
            with self.assertRaises(ValueError):
                ingest_file(path)

    def test_private_staging_copies_and_purges_file(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.txt"
            root = Path(directory) / "private-staging"
            source.write_text("Private source", encoding="utf-8")
            staging = LocalStaging(root)
            staged = staging.stage(source)
            self.assertEqual(staged.parent, root.resolve())
            self.assertEqual(staged.read_text(encoding="utf-8"), "Private source")
            staging.purge(staged)
            self.assertFalse(staged.exists())
            staged_bytes = staging.stage_bytes(".txt", b"private bytes")
            self.assertEqual(staged_bytes.read_bytes(), b"private bytes")
            staging.purge(staged_bytes)

    def test_legacy_office_formats_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            for suffix in (".doc", ".xls"):
                path = Path(directory) / f"legacy{suffix}"
                path.write_bytes(b"not supported")
                with self.assertRaisesRegex(ValueError, "Unsupported file type"):
                    ingest_file(path)

    def test_docx_and_xlsx_preserve_structured_values(self) -> None:
        from docx import Document
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            docx_path = Path(directory) / "guide.docx"
            document = Document()
            document.add_heading("Pricing", level=1)
            document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Plan"
            document.tables[0].rows[0].cells[1].text = "100 credits"
            document.save(docx_path)
            docx_text, _ = extract_file(docx_path)
            self.assertIn("Pricing", docx_text)
            self.assertIn("100 credits", docx_text)

            xlsx_path = Path(directory) / "pricing.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Plans"
            sheet.append(["Plan", "Units"])
            sheet.append(["Starter", 100])
            workbook.save(xlsx_path)
            xlsx_text, _ = extract_file(xlsx_path)
            self.assertIn("Sheet: Plans", xlsx_text)
            self.assertIn("Starter | 100", xlsx_text)

    def test_json_and_csv_preserve_structured_values(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            json_path = Path(directory) / "metadata.json"
            json_path.write_text('{"name":"Starter","credits":100}', encoding="utf-8")
            json_text, json_mime = extract_file(json_path)
            self.assertEqual(json_mime, "application/json")
            self.assertIn('"credits": 100', json_text)

            jsonl_path = Path(directory) / "events.jsonl"
            jsonl_path.write_text('{"name":"Starter"}\n\n{"name":"Pro","credits":200}\n', encoding="utf-8")
            jsonl_text, jsonl_mime = extract_file(jsonl_path)
            self.assertEqual(jsonl_mime, "application/x-ndjson")
            self.assertIn('"name": "Starter"', jsonl_text)
            self.assertIn('"credits": 200', jsonl_text)

            invalid_jsonl = Path(directory) / "invalid.jsonl"
            invalid_jsonl.write_text('{"valid":true}\nnot-json\n', encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Invalid JSONL source at line 2"):
                extract_file(invalid_jsonl)

            csv_path = Path(directory) / "catalog.csv"
            csv_path.write_text("name,credits\nStarter,100\n", encoding="utf-8")
            csv_text, csv_mime = extract_file(csv_path)
            self.assertEqual(csv_mime, "text/csv")
            self.assertIn("Starter | 100", csv_text)

    def test_zip_extracts_supported_files_with_file_locators(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "portfolio.zip"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("bundle/training/professional-profile.md", "# Professional profile\n\nSecure AI delivery.")
                archive.writestr("bundle/training/career-timeline.txt", "2026: Platform launch.")
                archive.writestr("bundle/training/catalog.csv", "name,credits\nStarter,100\n")
                archive.writestr("bundle/training/metadata.json", '{"name":"Starter"}')
                archive.writestr("bundle/training/events.jsonl", '{"name":"Starter"}\n{"credits":100}\n')
                archive.writestr("bundle/evaluation_do_not_ingest/questions.json", '{"question":"ignore"}')
                archive.writestr("bundle/README.md", "Do not ingest this README.")

            source = ingest_file(path)

            self.assertEqual(source.mime_type, "application/zip")
            self.assertIn("professional-profile.md", source.normalized_content)
            self.assertIn("career-timeline.txt", source.normalized_content)
            self.assertIn("catalog.csv", source.normalized_content)
            self.assertIn("metadata.json", source.normalized_content)
            self.assertIn("events.jsonl", source.normalized_content)
            self.assertNotIn("questions.json", source.normalized_content)
            self.assertNotIn("README.md", source.normalized_content)
            self.assertTrue(any(chunk.locator.startswith("bundle/training/professional-profile.md") for chunk in source.chunks))
            self.assertTrue(any(chunk.locator.startswith("bundle/training/career-timeline.txt") for chunk in source.chunks))

    def test_zip_rejects_unsupported_nested_and_unsafe_members(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            unsupported = Path(directory) / "unsupported.zip"
            with zipfile.ZipFile(unsupported, "w") as archive:
                archive.writestr("payload.exe", "not supported")
            with self.assertRaisesRegex(ValueError, "unsupported file type"):
                ingest_file(unsupported)

            nested = Path(directory) / "nested.zip"
            with zipfile.ZipFile(nested, "w") as archive:
                archive.writestr("nested.zip", "PK\\x05\\x06")
            with self.assertRaisesRegex(ValueError, "Nested ZIP"):
                ingest_file(nested)

            unsafe = Path(directory) / "unsafe.zip"
            with zipfile.ZipFile(unsafe, "w") as archive:
                archive.writestr("../outside.txt", "unsafe")
            with self.assertRaisesRegex(ValueError, "unsafe path"):
                ingest_file(unsafe)

    def test_zip_accepts_content_over_the_previous_200k_limit(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "large.zip"
            records = ''.join('{"value":"' + ('x' * 80) + '"}\n' for _ in range(3_000))
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("bundle/data/events.jsonl", records)

            source = ingest_file(path)

            self.assertGreater(len(source.normalized_content), 200_000)

    def test_pdf_locators_and_ocr_only_rejection(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "faq.pdf"
            path.write_bytes(b"%PDF-1.7 synthetic")
            page = Mock()
            page.extract_text.return_value = "Pricing is 100 credits."
            with patch("pypdf.PdfReader", return_value=Mock(pages=[page])):
                text, mime = extract_file(path)
            self.assertEqual(mime, "application/pdf")
            self.assertIn("Page 1:", text)
            empty_page = Mock()
            empty_page.extract_text.return_value = ""
            with patch("pypdf.PdfReader", return_value=Mock(pages=[empty_page])):
                with self.assertRaisesRegex(ValueError, "OCR-only"):
                    extract_file(path)

    def test_office_container_signatures_and_macros_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            fake_docx = Path(directory) / "fake.docx"
            fake_docx.write_bytes(b"not a zip")
            with self.assertRaisesRegex(ValueError, "signature"):
                extract_file(fake_docx)

            macro_docx = Path(directory) / "macro.docx"
            with zipfile.ZipFile(macro_docx, "w") as archive:
                archive.writestr("word/vbaProject.bin", b"macro")
            with self.assertRaisesRegex(ValueError, "Macro"):
                extract_file(macro_docx)

            unsafe_docx = Path(directory) / "unsafe.docx"
            with zipfile.ZipFile(unsafe_docx, "w") as archive:
                archive.writestr("../outside.txt", b"unsafe")
            with self.assertRaisesRegex(ValueError, "unsafe path"):
                extract_file(unsafe_docx)

            with patch("ingestion.pipeline.MAX_CONTAINER_UNCOMPRESSED_BYTES", 1):
                with self.assertRaisesRegex(ValueError, "extraction limits"):
                    extract_file(macro_docx)

    def test_website_ingestion_blocks_private_dns_and_redirects(self) -> None:
        with patch("ingestion.pipeline.socket.getaddrinfo", return_value=[(None, None, None, None, ("127.0.0.1", 443))]), patch("requests.get") as request:
            with self.assertRaisesRegex(ValueError, "Private"):
                fetch_authorized_https("https://approved.example/faq", {"approved.example"})
            request.assert_not_called()

        response = Mock(status_code=200, encoding="utf-8")
        response.headers = {"content-type": "text/html; charset=utf-8"}
        response.iter_content.return_value = [b"<html><body>Approved FAQ</body></html>"]
        response.raise_for_status.return_value = None
        with patch("ingestion.pipeline.socket.getaddrinfo", return_value=[(None, None, None, None, ("93.184.216.34", 443))]), patch("requests.get", return_value=response):
            text, mime = fetch_authorized_https("https://approved.example/faq", {"approved.example"})
        self.assertEqual(mime, "text/html")
        self.assertIn("Approved FAQ", text)

        redirect = Mock(status_code=302, encoding="utf-8")
        redirect.raise_for_status.return_value = None
        with patch("ingestion.pipeline.socket.getaddrinfo", return_value=[(None, None, None, None, ("93.184.216.34", 443))]), patch("requests.get", return_value=redirect):
            with self.assertRaisesRegex(ValueError, "Redirects"):
                fetch_authorized_https("https://approved.example/redirect", {"approved.example"})

        response = Mock(status_code=200, encoding="utf-8")
        response.headers = {"content-type": "text/html"}
        response.iter_content.return_value = [b"<html><body>Rebound</body></html>"]
        response.raise_for_status.return_value = None
        public = (None, None, None, None, ("93.184.216.34", 443))
        private = (None, None, None, None, ("10.0.0.8", 443))
        with patch("ingestion.pipeline.socket.getaddrinfo", side_effect=[[public], [private]]), patch("requests.get", return_value=response) as request:
            with self.assertRaisesRegex(ValueError, "Private"):
                fetch_authorized_https("https://approved.example/rebound", {"approved.example"})
            request.assert_called_once()

    def test_url_ingestion_keeps_approved_snapshot_url_and_sanitizes_content(self) -> None:
        response = Mock(status_code=200, encoding="utf-8")
        response.headers = {"content-type": "text/html"}
        response.iter_content.return_value = [b"<html><body>Pricing is 42 credits. Email alice@example.com</body></html>"]
        response.raise_for_status.return_value = None
        with patch("ingestion.pipeline.socket.getaddrinfo", return_value=[(None, None, None, None, ("93.184.216.34", 443))]), patch("requests.get", return_value=response):
            source = ingest_url("Approved FAQ", "https://approved.example/faq", {"approved.example"})
        self.assertEqual(source.source_url, "https://approved.example/faq")
        self.assertIn("Pricing is 42 credits.", source.normalized_content)
        self.assertNotIn("alice@example.com", source.normalized_content)

        binary = Mock(status_code=200, encoding="utf-8", headers={"content-type": "application/zip"})
        binary.raise_for_status.return_value = None
        with patch("ingestion.pipeline.socket.getaddrinfo", return_value=[(None, None, None, None, ("93.184.216.34", 443))]), patch("requests.get", return_value=binary):
            with self.assertRaisesRegex(ValueError, "text content"):
                fetch_authorized_https("https://approved.example/archive", {"approved.example"})

    def test_worker_payload_is_sanitized_and_submission_claims_then_publishes(self) -> None:
        source = ingest_text("FAQ", "Email alice@example.com.\n\nPricing is 42 credits.", "text/plain")
        payload = result_payload(source)
        self.assertNotIn("alice@example.com", payload["normalizedContent"])
        self.assertEqual(payload["chunks"][0]["contentHash"], source.chunks[0].content_hash)
        self.assertNotIn("content_hash", payload["chunks"][0])
        session = Mock()
        claim_response = Mock()
        result_response = Mock()
        result_response.json.return_value = {"job": {"status": "completed"}}
        session.post.side_effect = [claim_response, result_response]
        output = submit_result("http://localhost:4000/", "job-1", "worker-secret", source, session)
        self.assertEqual(output["job"]["status"], "completed")
        self.assertEqual(session.post.call_count, 2)
        self.assertEqual(session.post.call_args_list[0].args[0], "http://localhost:4000/api/v1/internal/ingestion-jobs/job-1/claim")
        self.assertEqual(session.post.call_args_list[1].kwargs["json"]["contentHash"], source.content_hash)

    def test_worker_fetches_api_file_input_after_claim(self) -> None:
        session = Mock()
        response = Mock()
        response.json.return_value = {"fileName": "job-1.txt", "mimeType": "text/plain", "contentBase64": "RkFRIGRhdGE="}
        session.get.return_value = response
        file_name, mime_type, content = fetch_job_input("http://localhost:4000/", "job-1", "worker-secret", session)
        self.assertEqual((file_name, mime_type, content), ("job-1.txt", "text/plain", b"FAQ data"))
        self.assertEqual(session.get.call_args.args[0], "http://localhost:4000/api/v1/internal/ingestion-jobs/job-1/input")

    def test_worker_claims_next_job_or_waits_when_queue_is_empty(self) -> None:
        session = Mock()
        response = Mock()
        response.json.return_value = {"job": {"id": "job-1"}}
        session.post.return_value = response
        self.assertEqual(claim_next_job("http://localhost:4000", "worker-secret", session), "job-1")
        self.assertEqual(session.post.call_args.args[0], "http://localhost:4000/api/v1/internal/ingestion-jobs/claim-next")

        response.json.return_value = {"job": None}
        self.assertIsNone(claim_next_job("http://localhost:4000", "worker-secret", session))

    def test_document_worker_renews_its_lease(self) -> None:
        session = Mock()
        response = Mock()
        response.json.return_value = {"job": {"id": "job-1", "status": "processing"}}
        session.post.return_value = response
        renewed = renew_document_job("http://localhost:4000", "job-1", "tenant-1", "worker-secret", session=session)
        self.assertEqual(renewed["status"], "processing")
        self.assertEqual(session.post.call_args.kwargs["json"], {"tenantId": "tenant-1", "jobId": "job-1", "leaseMs": 120_000})

    def test_worker_submits_without_reclaiming_an_already_claimed_upload(self) -> None:
        source = ingest_text("FAQ", "Approved content", "text/plain")
        session = Mock()
        response = Mock()
        response.json.return_value = {"job": {"status": "completed"}}
        session.post.return_value = response
        submit_result("http://localhost:4000", "job-1", "worker-secret", source, session, claim=False)
        self.assertEqual(session.post.call_count, 1)
        self.assertIn("/api/v1/internal/ingestion-jobs/job-1/result", session.post.call_args.args[0])

    def test_worker_reports_extraction_failure(self) -> None:
        session = Mock()
        response = Mock()
        response.json.return_value = {"job": {"status": "failed"}}
        session.post.return_value = response
        output = submit_failure("http://localhost:4000", "job-1", "worker-secret", session=session)
        self.assertEqual(output["job"]["status"], "failed")
        self.assertEqual(session.post.call_args.kwargs["json"], {"errorCode": "INGESTION_FAILED"})

    def test_worker_preserves_safe_api_error_code(self) -> None:
        source = ingest_text("FAQ", "Approved content", "text/plain")
        session = Mock()
        response = Mock()
        response.raise_for_status.side_effect = requests.HTTPError("422 Client Error")
        response.json.return_value = {"error": {"code": "CHUNK_HASH_MISMATCH"}}
        session.post.return_value = response
        with self.assertRaisesRegex(requests.HTTPError, "CHUNK_HASH_MISMATCH"):
            submit_result("http://localhost:4000", "job-1", "worker-secret", source, session, claim=False)

    def test_worker_claims_before_extracting_path_or_url_inputs(self) -> None:
        source = ingest_text("FAQ", "Approved content", "text/plain")
        events: list[str] = []

        def claim(*_args: object, **_kwargs: object) -> None:
            events.append("claim")

        def extract(_path: Path) -> object:
            events.append("extract")
            return source

        def submit(*_args: object, **kwargs: object) -> None:
            events.append(f"submit:{kwargs['claim']}")

        with patch.object(sys, "argv", ["worker", "--path", "faq.txt", "--api-base", "http://localhost:4000", "--job-id", "job-1", "--worker-secret", "secret"]), patch("ingestion.worker.claim_job", side_effect=claim), patch("ingestion.worker.ingest_file", side_effect=extract), patch("ingestion.worker.submit_result", side_effect=submit):
            worker_main()

        self.assertEqual(events, ["claim", "extract", "submit:False"])


if __name__ == "__main__":
    unittest.main()
