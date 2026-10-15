from __future__ import annotations

import argparse
import base64
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any
from unittest.mock import Mock, patch

import requests
from docx import Document
from openpyxl import Workbook

from ingestion.pipeline import ingest_url
from ingestion.worker import claim_job, result_payload, submit_result


def request_json(client: requests.Session, method: str, base: str, path: str, *, expected: set[int], headers: dict[str, str] | None = None, body: dict[str, Any] | None = None) -> dict[str, Any]:
    response = client.request(method, f"{base.rstrip('/')}{path}", headers=headers, json=body, timeout=15)
    if response.status_code not in expected:
        raise AssertionError(f"{method} {path} returned {response.status_code}: {response.text}")
    return response.json() if response.content else {}


def pdf_bytes(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
    return bytes(output)


def write_fixtures(root: Path) -> list[tuple[str, str, str, bytes]]:
    txt = root / "accepted.txt"
    txt.write_text("TXT approved price is 10 credits.\n\nContact alice@example.com.", encoding="utf-8")

    docx = root / "accepted.docx"
    document = Document()
    document.add_heading("DOCX pricing", level=1)
    document.add_paragraph("DOCX approved price is 20 credits.")
    document.save(docx)

    xlsx = root / "accepted.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Plans"
    sheet.append(["Plan", "Price"])
    sheet.append(["XLSX approved", 30])
    workbook.save(xlsx)

    pdf = root / "accepted.pdf"
    pdf.write_bytes(pdf_bytes("PDF approved price is 40 credits."))

    return [
        ("Accepted TXT", "accepted.txt", "text/plain", txt.read_bytes()),
        ("Accepted DOCX", "accepted.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx.read_bytes()),
        ("Accepted XLSX", "accepted.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx.read_bytes()),
        ("Accepted PDF", "accepted.pdf", "application/pdf", pdf.read_bytes()),
    ]


def run_worker(api_base: str, job_id: str, worker_secret: str, staging_root: Path) -> None:
    command = [sys.executable, "-m", "ingestion.worker", "--staging-root", str(staging_root), "--api-base", api_base, "--job-id", job_id, "--worker-secret", worker_secret]
    result = subprocess.run(command, cwd=Path(__file__).resolve().parent, capture_output=True, text=True, check=False)
    if result.returncode != 0:
        raise AssertionError(f"worker failed for {job_id}: {result.stdout}\n{result.stderr}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Run the deterministic local five-source ingestion/API acceptance journey.")
    parser.add_argument("--api-base", required=True)
    parser.add_argument("--worker-secret", required=True)
    args = parser.parse_args()

    client = requests.Session()
    platform_credentials = {"email": "acceptance-platform@example.test", "password": "Platform1Password"}
    platform = request_json(client, "POST", args.api_base, "/api/v1/auth/register-platform", expected={201}, body=platform_credentials)
    tenant = request_json(client, "POST", args.api_base, "/api/v1/admin/tenants", expected={201}, headers={"Authorization": f"Bearer {platform['token']}"}, body={"name": "Ingestion Acceptance", "slug": "ingestion-acceptance", "adminEmail": "acceptance-admin@example.test", "adminPassword": "TenantA1Password"})
    admin = request_json(client, "POST", args.api_base, "/api/v1/auth/login", expected={200}, body={"email": "acceptance-admin@example.test", "password": "TenantA1Password"})
    admin_headers = {"Authorization": f"Bearer {admin['token']}"}
    request_json(client, "POST", args.api_base, "/api/v1/admin/domains", expected={201}, headers=admin_headers, body={"hostname": "approved.example"})

    with tempfile.TemporaryDirectory(prefix="agentic-ingestion-fixtures-") as directory:
        fixture_root = Path(directory)
        worker_staging = fixture_root / "worker-staging"
        fixture_sources = write_fixtures(fixture_root)
        published: list[str] = []
        for title, file_name, mime_type, content in fixture_sources:
            queued = request_json(client, "POST", args.api_base, "/api/v1/admin/sources/upload", expected={202}, headers=admin_headers, body={"title": title, "fileName": file_name, "mimeType": mime_type, "audience": "public", "contentBase64": base64.b64encode(content).decode("ascii")})
            run_worker(args.api_base, queued["job"]["id"], args.worker_secret, worker_staging)
            completed = request_json(client, "GET", args.api_base, f"/api/v1/admin/ingestion-jobs/{queued['job']['id']}", expected={200}, headers=admin_headers)
            if completed["job"]["status"] != "completed" or not completed["job"]["sourceId"]:
                raise AssertionError(f"file job did not complete: {completed}")
            source = request_json(client, "POST", args.api_base, f"/api/v1/admin/sources/{completed['job']['sourceId']}/publish", expected={200}, headers=admin_headers)
            published.append(source["source"]["id"])

        url_job = request_json(client, "POST", args.api_base, "/api/v1/admin/ingestion-jobs", expected={202}, headers=admin_headers, body={"title": "Accepted HTTPS snapshot", "mimeType": "text/html", "sourceUrl": "https://approved.example/faq", "audience": "public"})
        response = Mock(status_code=200, encoding="utf-8", headers={"content-type": "text/html"})
        response.iter_content.return_value = [b"<html><body>URL approved price is 50 credits.</body></html>"]
        response.raise_for_status.return_value = None
        public_address = (None, None, None, None, ("93.184.216.34", 443))
        with patch("ingestion.pipeline.socket.getaddrinfo", return_value=[public_address]), patch("requests.get", return_value=response):
            url_source = ingest_url("Accepted HTTPS snapshot", "https://approved.example/faq", {"approved.example"})
        claim_job(args.api_base, url_job["job"]["id"], args.worker_secret, client)
        submit_result(args.api_base, url_job["job"]["id"], args.worker_secret, url_source, client, claim=False)
        url_completed = request_json(client, "GET", args.api_base, f"/api/v1/admin/ingestion-jobs/{url_job['job']['id']}", expected={200}, headers=admin_headers)
        if url_completed["job"]["status"] != "completed" or not url_completed["job"]["sourceId"]:
            raise AssertionError(f"URL job did not complete: {url_completed}")
        request_json(client, "POST", args.api_base, f"/api/v1/admin/sources/{url_completed['job']['sourceId']}/publish", expected={200}, headers=admin_headers)
        published.append(url_completed["job"]["sourceId"])

        sources = request_json(client, "GET", args.api_base, "/api/v1/admin/sources", expected={200}, headers=admin_headers)
        active_titles = {source["title"] for source in sources["items"] if source["status"] == "active"}
        expected_titles = {title for title, *_ in fixture_sources} | {"Accepted HTTPS snapshot"}
        if not expected_titles.issubset(active_titles) or len(published) != 5:
            raise AssertionError(f"not all five sources were published: {active_titles}")

        channel_headers = {**admin_headers, "content-type": "application/json"}
        request_json(client, "PATCH", args.api_base, "/api/v1/admin/channel", expected={200}, headers=channel_headers, body={"publicOrigin": "http://localhost:3000", "publicChatEnabled": True})
        public_headers = {"Origin": "http://localhost:3000"}
        session = request_json(client, "POST", args.api_base, "/api/v1/chat/sessions", expected={201}, headers=public_headers, body={"tenantSlug": "ingestion-acceptance"})
        answer = request_json(client, "POST", args.api_base, "/api/v1/chat/messages", expected={200}, headers=public_headers, body={"sessionId": session["sessionId"], "message": "What is the XLSX plan?", "clientMessageId": "ingestion-acceptance-chat"})
        if not answer["grounded"] or not answer["citations"]:
            raise AssertionError(f"published ingestion sources did not ground chat: {answer}")

    print(json.dumps({"sourcesPublished": 5, "fileTypes": ["TXT", "DOCX", "XLSX", "PDF"], "urlSnapshot": True, "groundedChat": True}, sort_keys=True))


if __name__ == "__main__":
    main()
